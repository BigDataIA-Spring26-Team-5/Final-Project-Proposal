"""
Convert data_sources_revised.md to DOCX
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re

doc = Document()

# Style setup
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# Read the markdown
with open("data_sources_revised.md", "r") as f:
    content = f.read()

lines = content.split("\n")
i = 0
in_table = False
table_rows = []
in_code_block = False
code_lines = []

def add_table(doc, rows):
    if not rows:
        return
    # Parse markdown table rows
    parsed = []
    for row in rows:
        cells = [c.strip() for c in row.strip("|").split("|")]
        parsed.append(cells)

    if len(parsed) < 2:
        return

    # Skip separator row (row with ---)
    header = parsed[0]
    data = [r for r in parsed[1:] if not all(set(c.strip()) <= set("-: ") for c in r)]

    if not data:
        return

    num_cols = len(header)
    table = doc.add_table(rows=1 + len(data), cols=num_cols, style='Light Grid Accent 1')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    for j, cell_text in enumerate(header):
        if j < num_cols:
            cell = table.rows[0].cells[j]
            cell.text = cell_text.strip().replace("**", "")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(10)

    # Data rows
    for i, row_data in enumerate(data):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.rows[i + 1].cells[j]
                cell.text = cell_text.strip().replace("**", "").replace("`", "")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

while i < len(lines):
    line = lines[i]

    # Code blocks
    if line.strip().startswith("```"):
        if in_code_block:
            # End code block - add as formatted text
            code_text = "\n".join(code_lines)
            p = doc.add_paragraph()
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            code_lines = []
            in_code_block = False
        else:
            in_code_block = True
            code_lines = []
        i += 1
        continue

    if in_code_block:
        code_lines.append(line)
        i += 1
        continue

    # Tables
    if "|" in line and line.strip().startswith("|"):
        if not in_table:
            in_table = True
            table_rows = []
        table_rows.append(line)
        i += 1
        continue
    elif in_table:
        add_table(doc, table_rows)
        table_rows = []
        in_table = False

    # Headers
    if line.startswith("# ") and not line.startswith("##"):
        doc.add_heading(line[2:].strip(), level=0)
    elif line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
    elif line.startswith("---"):
        doc.add_paragraph("─" * 60)
    elif line.strip() == "":
        pass  # Skip blank lines
    elif line.startswith("- **") or line.startswith("  - **"):
        # Bold list items
        p = doc.add_paragraph(style='List Bullet')
        # Parse bold markers
        text = line.strip().lstrip("- ")
        parts = re.split(r'\*\*', text)
        for idx, part in enumerate(parts):
            if part:
                run = p.add_run(part.replace("`", ""))
                run.font.size = Pt(11)
                if idx % 2 == 1:  # Bold parts
                    run.bold = True
    elif line.startswith("- ") or line.startswith("  - "):
        p = doc.add_paragraph(line.strip().lstrip("- ").replace("**", "").replace("`", ""), style='List Bullet')
    else:
        # Regular paragraph - handle bold and code markers
        text = line.strip()
        if text:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2].replace("`", ""))
                    run.bold = True
                else:
                    run = p.add_run(part.replace("`", ""))
                run.font.size = Pt(11)

    i += 1

# Handle any remaining table
if in_table and table_rows:
    add_table(doc, table_rows)

output_path = "data_sources_revised.docx"
doc.save(output_path)
print(f"Saved: {output_path}")
